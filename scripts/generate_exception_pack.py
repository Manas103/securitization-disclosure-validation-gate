"""One command regenerates docs/exception_pack.xlsx: a flat, pivot-ready
Excel workbook a non-SQL reviewer can open directly, built from a run of the
gate against the 60-seeded-violation demo load.

    python scripts/generate_exception_pack.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from app import db
from app.compiler import compile_assertions
from app.engine import run_all
from app.spec_loader import load_spec
from app.violations import inject_seeded_violations

OUT_PATH = Path(__file__).resolve().parents[1] / "docs" / "exception_pack.xlsx"


def main() -> int:
    attributes = load_spec()
    assertions = compile_assertions(attributes)
    conn = db.connect()

    files, _ = inject_seeded_violations(attributes, seed=777, n_records=120)
    for file_type, records in files.items():
        db.load_file(conn, file_type, records)
    violations = run_all(conn, assertions)

    wb = openpyxl.Workbook()

    assertion_sheet = wb.active
    assertion_sheet.title = "Assertion Catalog"
    assertion_sheet.append(["assertion_id", "attribute_id", "attribute_name", "assertion_type", "description"])
    for a in assertions:
        assertion_sheet.append([a.assertion_id, a.attribute_id, a.attribute_name, a.assertion_type, a.description])

    exceptions_sheet = wb.create_sheet("Exceptions Detail")
    exceptions_sheet.append(["file_type", "record_id", "attribute_id", "attribute_name", "assertion_type", "offending_value"])
    for v in violations:
        exceptions_sheet.append([v.file_type, v.record_id, v.attribute_id, v.attribute_name, v.assertion_type, v.value])

    summary_sheet = wb.create_sheet("Summary By File")
    summary_sheet.append(["file_type", "records_loaded", "exceptions_found"])
    counts = {}
    for file_type, records in files.items():
        counts[file_type] = sum(1 for v in violations if v.file_type == file_type)
    for file_type, records in files.items():
        summary_sheet.append([file_type, len(records), counts[file_type]])

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH} ({OUT_PATH.stat().st_size} bytes)")
    print(f"Assertion Catalog: {len(assertions)} rows")
    print(f"Exceptions Detail: {len(violations)} rows")
    print(f"Summary By File: {len(files)} rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
